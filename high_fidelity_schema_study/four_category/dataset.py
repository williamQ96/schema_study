"""Deterministic, evidence-linked dataset adapter for the four-category study.

This module deliberately makes no language-model calls. Observations are kept
separate from declarations and from conservative name-based hints.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
from pathlib import PureWindowsPath
import csv
from importlib import metadata as importlib_metadata
from collections import defaultdict
import zipfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any


SCHEMA_VERSION = "dataset-four-category/v1"
CATEGORIES = {"structure", "encoding", "value", "syntax"}


def _canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _hash_file(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
            size += len(chunk)
    return digest.hexdigest(), size


def _json_value(value: Any) -> Any:
    if isinstance(value, bytes):
        return value.decode("utf-8", "replace")
    if hasattr(value, "tolist"):
        return _json_value(value.tolist())
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return "NaN" if math.isnan(value) else ("Infinity" if value > 0 else "-Infinity")
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(v) for v in value]
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _xlsx_raw_cells(workbook: Any, worksheet: Any, max_row: int) -> dict[str, dict]:
    """Retain original worksheet XML tokens for the bounded row prefix."""
    tokens = {}
    stack = []
    with workbook._archive.open(worksheet._worksheet_path) as stream:
        for event, elem in ET.iterparse(stream, events=("start", "end")):
            if event == "start":
                stack.append(elem)
                continue
            local = elem.tag.rsplit("}", 1)[-1]
            if local == "c":
                coord = elem.attrib.get("r", "")
                try:
                    row = int(re.search(r"\d+", coord).group())
                except (AttributeError, ValueError):
                    row = 0
                if row <= max_row:
                    values = {child.tag.rsplit("}", 1)[-1]: child.text for child in elem}
                    inline = "".join(node.text or "" for node in elem.iter() if node.tag.rsplit("}", 1)[-1] == "t")
                    tokens[coord] = {"xml_cell_type": elem.attrib.get("t"),
                                     "xml_style_index": elem.attrib.get("s"),
                                     "xml_value_token": values.get("v"),
                                     "xml_formula_token": values.get("f"),
                                     "xml_inline_text": inline or None}
                if len(stack) > 1:
                    stack[-2].remove(elem)
            elif local == "row":
                row = int(elem.attrib.get("r", "0") or 0)
                if row > max_row:
                    break
                if len(stack) > 1:
                    stack[-2].remove(elem)
            stack.pop()
    return tokens


def _taxonomy_data(taxonomy: Any) -> Any:
    if taxonomy is None:
        path = Path(__file__).resolve().parents[1] / "templates" / "four_category_taxonomy_v1.json"
        return json.loads(path.read_text(encoding="utf-8"))
    if isinstance(taxonomy, (str, Path)):
        return json.loads(Path(taxonomy).read_text(encoding="utf-8"))
    return taxonomy


def _relative(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError as exc:
        raise ValueError(f"source {path} is outside root {root}") from exc


def _zarr_inventory(path: Path) -> tuple[list[dict], bytes]:
    """Hash Zarr-v2 metadata members only; chunk payloads are intentionally excluded."""
    suffixes = {".zgroup", ".zarray", ".zattrs", ".zmetadata"}
    members = []
    for member in path.rglob("*"):
        if member.is_file() and member.name in suffixes:
            if not member.resolve().is_relative_to(path.resolve()):
                raise ValueError(f"Zarr metadata member escapes store root: {member.name}")
            digest, size = _hash_file(member)
            members.append({"path": member.relative_to(path).as_posix(), "sha256": digest, "bytes": size})
    if not members:
        raise ValueError("zarr-v2 store has no .zgroup, .zarray, .zattrs, or .zmetadata metadata")
    members.sort(key=lambda item: item["path"])
    return members, _canonical(members)


def _dependency_versions(fmt: str) -> dict:
    names = {"xlsx": ["openpyxl"], "arff": ["scipy"], "hdf5": ["h5py"],
             "netcdf": ["scipy", "h5py"], "parquet": ["pyarrow"]}.get(fmt, [])
    versions = {}
    for name in names:
        try:
            versions[name] = importlib_metadata.version(name)
        except importlib_metadata.PackageNotFoundError:
            versions[name] = None
    return versions


def _unit_hint(name: str) -> str | None:
    # A name suffix is a weak inference only. In particular it never makes an
    # observed value or an encoding declaration.
    m = re.search(r"(?:_|\[)(mm|cm|km|kg|g|m|s|ms|c|f|k|pct|percent|mg_l|m_s)$", name.lower().rstrip("]"))
    return m.group(1) if m else None


def _fact(bag: dict, subject: str, predicate: str, value: Any, category: str,
          basis: str, source_path: str, locator: dict, raw: Any,
          coverage: dict) -> None:
    eid = f"{bag['dataset_id']}:e{len(bag['evidence']) + 1:05d}"
    bag["evidence"].append({"evidence_id": eid, "source_path": source_path,
                            "locator": locator, "raw_value": raw, "coverage": coverage})
    fid = f"{bag['dataset_id']}:f{len(bag['facts']) + 1:05d}"
    categories = [category] if isinstance(category, str) else sorted(set(category))
    bag["facts"].append({"fact_id": fid, "subject_id": subject,
                         "predicate": predicate, "value": value,
                         "categories": categories, "basis": basis,
                         "evidence_ids": [eid], "coverage": coverage})


def _read_delimited(path: Path, fmt: str, limit: int) -> tuple[list[list[str]], list[str], dict]:
    delimiter = "\t" if fmt == "tsv" else ","
    sample = ""
    dialect = csv.excel_tab if fmt == "tsv" else csv.excel
    dialect_basis = "tsv_format_default" if fmt == "tsv" else "csv_default"
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(8192)
    except (OSError, UnicodeDecodeError):
        pass
    if fmt == "csv":
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            delimiter = dialect.delimiter
            dialect_basis = "bounded_sample_sniff"
        except (OSError, UnicodeDecodeError, csv.Error):
            pass
    rows: list[list[str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, dialect=dialect)
        total = 0
        for row in reader:
            total += 1
            if len(rows) < max(1, limit + 1):
                rows.append(row)
    if not rows:
        return [], [], {"mode": "full", "rows_read": 0, "delimiter": delimiter,
                        "quotechar": dialect.quotechar, "total_rows": 0,
                        "quotechar_observed": bool(dialect.quotechar and dialect.quotechar in sample),
                        "delimiter_observed": bool(delimiter and delimiter in sample),
                        "doublequote": dialect.doublequote, "escapechar": dialect.escapechar,
                        "quoting": dialect.quoting, "dialect_basis": dialect_basis}
    header = rows[0]
    selected = rows[1:]
    sample_mode = total - 1 > len(selected)
    return selected, header, {"mode": "sample" if sample_mode else "full",
                              "sample_limit": limit, "rows_read": len(selected), "total_rows": max(0,total-1), "header_rows": 1,
                              "delimiter": delimiter, "quotechar": dialect.quotechar,
                              "quotechar_observed": bool(dialect.quotechar and dialect.quotechar in sample),
                              "delimiter_observed": bool(delimiter and delimiter in sample),
                              "doublequote": dialect.doublequote, "escapechar": dialect.escapechar,
                              "quoting": dialect.quoting, "dialect_basis": dialect_basis}


def _parse(path: Path, rel: str, fmt: str, limit: int, bag: dict) -> tuple[str, list[str]]:
    issues: list[str] = []
    if fmt in {"csv", "tsv"}:
        rows, headers, cov = _read_delimited(path, fmt, limit)
        bag["parser"]["delimited_parser"] = {key: cov.get(key) for key in
                                              ("delimiter", "quotechar", "doublequote", "escapechar", "quoting", "dialect_basis")}
        _fact(bag, "dataset", "serialization_format", fmt, "encoding", "observed", rel,
              {"kind": "validated_serialization", "format": fmt}, fmt,
              {"mode": "full", "rows_validated": cov.get("total_rows", 0)})
        for i, h in enumerate(headers):
            loc = {"kind": "delimited_header", "row": 1, "column": i + 1}
            _fact(bag, f"column:{i+1}", "name", h, "structure", "observed", rel, loc, h, {"mode": "full", "row": 1})
            _fact(bag, f"column:{i+1}", "position", i + 1, "syntax", "observed", rel, loc, h, {"mode": "full", "row": 1})
        delimiter = cov.get("delimiter", "\t" if fmt == "tsv" else ",")
        if cov.get("delimiter_observed"):
            _fact(bag, "dataset", "delimiter", {"\t": "tab", ",": "comma", ";": "semicolon", "|": "pipe"}.get(delimiter, delimiter), "syntax", "observed", rel,
                  {"kind": "delimited_header_syntax", "row": 1}, delimiter, {"mode": "sample", "sample_limit_bytes": 8192})
        if cov.get("quotechar_observed"):
            _fact(bag, "dataset", "quote_character", cov.get("quotechar"), "syntax", "observed", rel,
                  {"kind": "delimited_header_syntax", "row": 1}, cov.get("quotechar"), {"mode": "sample", "sample_limit_bytes": 8192})
        for ri, row in enumerate(rows, 2):
            for ci, val in enumerate(row):
                _fact(bag, f"column:{ci+1}", "observed_lexical_value", val, "syntax", "observed", rel,
                      {"kind": "delimited_cell", "row": ri, "column": ci+1}, val,
                      {**cov, "row": ri})
        for ci, name in enumerate(headers):
            unit = _unit_hint(name)
            if unit:
                _fact(bag, f"column:{ci+1}", "unit_hint", unit, "value", "inferred", rel,
                      {"kind": "delimited_header", "row": 1, "column": ci+1}, name, {"mode": "full", "row": 1})
        return fmt, issues
    if fmt in {"json", "jsonl"}:
        if fmt == "json" and path.stat().st_size > 8 * 1024 * 1024:
            return fmt, ["unsupported: single-document JSON exceeds 8 MiB bounded parse limit"]
        # Keep JSON number/string token spellings where possible by parsing
        # numeric tokens as strings. This avoids lexical loss in evidence.
        try:
            total_records = 1
            if fmt == "jsonl":
                payload = []
                total_records = 0
                with path.open(encoding="utf-8-sig") as stream:
                    for line_no in range(1, 2**63):
                        line = stream.readline(1024 * 1024 + 1)
                        if not line: break
                        if len(line) > 1024 * 1024:
                            raise ValueError(f"JSONL line {line_no} exceeds 1 MiB bounded line limit")
                        if line.strip():
                            total_records += 1
                            obj = json.loads(line, parse_int=str, parse_float=str)
                            if len(payload) < max(1, limit): payload.append((line_no, obj))
            else:
                payload = [(1, json.loads(path.read_text(encoding="utf-8-sig"), parse_int=str, parse_float=str))]
        except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError(f"malformed {fmt}: {exc}") from exc
        source_fact_start = len(bag["facts"])
        _fact(bag, "dataset", "serialization_format", fmt, "encoding", "observed", rel,
              {"kind": "validated_serialization", "format": fmt}, fmt,
              {"mode": "full", "records_validated": total_records if fmt == "jsonl" else 1})
        count = 0
        truncated = fmt == "jsonl" and total_records > len(payload)
        def visit(obj: Any, ptr: str, line: int, depth: int = 0) -> None:
            nonlocal count, truncated
            if depth > 32 or count >= 10000:
                truncated = True
                return
            count += 1
            if isinstance(obj, dict):
                for property_index, (key, val) in enumerate(obj.items(), 1):
                    if count >= 10000:
                        truncated = True
                        break
                    child = ptr + "/" + str(key).replace("~", "~0").replace("/", "~1")
                    _fact(bag, child or "/", "property", key, "structure", "observed", rel,
                          {"kind": "json_pointer", "pointer": child, "line": line}, key,
                          {"mode": "sample" if fmt == "jsonl" else "full", "line": line})
                    _fact(bag, child or "/", "property_lexical_position", property_index, "syntax", "observed", rel,
                          {"kind": "json_pointer", "pointer": ptr or "$", "member": key, "line": line}, key,
                          {"mode": "sample" if fmt == "jsonl" else "full", "line": line})
                    visit(val, child, line, depth+1)
            elif isinstance(obj, list):
                array_coverage = {"mode": "sample" if len(obj) > limit else "full", "line": line,
                                 "array_length": len(obj), "sample_limit": limit}
                if len(obj) > limit: truncated = True
                _fact(bag, ptr or "$", "array_length", len(obj), "structure", "observed", rel,
                      {"kind": "json_pointer", "pointer": ptr, "line": line}, len(obj), array_coverage)
                for i, val in enumerate(obj[:limit]): visit(val, f"{ptr}/{i}", line, depth+1)
            else:
                _fact(bag, ptr or "$", "observed_lexical_value", obj, "syntax", "observed", rel,
                      {"kind": "json_pointer", "pointer": ptr or "$", "line": line}, obj,
                      {"mode": "sample" if fmt == "jsonl" else "full", "line": line})
        for line, obj in payload: visit(obj, "$", line)
        if truncated:
            for fact in bag["facts"][source_fact_start:]:
                if fact["coverage"]["mode"] == "full": fact["coverage"]["mode"] = "sample"
            for evidence in bag["evidence"]:
                if evidence["evidence_id"] in {eid for fact in bag["facts"][source_fact_start:] for eid in fact["evidence_ids"]} and evidence["coverage"]["mode"] == "full":
                    evidence["coverage"]["mode"] = "sample"
        return fmt, issues
    if fmt in {"xml", "xsd"}:
        try:
            # Parse through EOF to validate the document while emitting only a
            # bounded prefix of element observations.
            stack = []
            element_nodes = []
            occurrences = defaultdict(int)
            seen = 0
            _fact(bag, "dataset", "serialization_format", fmt, "encoding", "observed", rel,
                  {"kind": "validated_serialization", "format": fmt}, fmt,
                  {"mode": "full", "validation": "complete_document"})
            for event, elem in ET.iterparse(path, events=("start", "end")):
                if event == "start":
                    tag = elem.tag.rsplit("}", 1)[-1]
                    stack.append(tag)
                    element_nodes.append(elem)
                    if len(stack) > 1024: raise ValueError("XML nesting exceeds 1024 level parser limit")
                    if seen < limit:
                        p = "/" + "/".join(stack)
                        occurrences[p] += 1
                        occurrence = occurrences[p]
                        _fact(bag, p, "element", tag, "structure", "observed", rel,
                              {"kind": "xml_element", "path": p, "occurrence": occurrence}, tag,
                              {"mode": "sample", "element_limit": limit})
                        if elem.tag.startswith("{"):
                            namespace = elem.tag[1:].split("}", 1)[0]
                            _fact(bag, p, "namespace", namespace, "syntax", "observed", rel,
                                  {"kind": "xml_namespace", "path": p, "occurrence": occurrence}, elem.tag, {"mode": "sample", "element_limit": limit})
                        for key, val in elem.attrib.items():
                            _fact(bag, p, "observed_lexical_attribute", {"name": key, "value": val}, "syntax", "observed", rel,
                                  {"kind": "xml_attribute", "path": p, "name": key}, val,
                                  {"mode": "sample", "element_limit": limit})
                    seen += 1
                else:
                    stack.pop()
                    finished = element_nodes.pop()
                    if element_nodes:
                        element_nodes[-1].remove(finished)
                    if seen <= limit and elem.text:
                        text = elem.text
                        _fact(bag, "/" + "/".join(stack + [elem.tag.rsplit("}",1)[-1]]), "observed_lexical_value", text,
                              "syntax", "observed", rel, {"kind": "xml_text", "element_path": "/" + "/".join(stack + [elem.tag.rsplit("}",1)[-1]]), "occurrence": occurrences.get("/" + "/".join(stack + [elem.tag.rsplit("}",1)[-1]]), 1)}, text,
                              {"mode": "sample", "element_limit": limit})
                    elem.clear()
        except (ET.ParseError, OSError) as exc:
            raise ValueError(f"malformed {fmt}: {exc}") from exc
        return fmt, issues
    if fmt == "xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=False)
        except ImportError:
            return fmt, ["dependency unavailable: openpyxl is required for xlsx"]
        _fact(bag, "dataset", "serialization_format", fmt, "encoding", "observed", rel,
              {"kind": "validated_serialization", "format": fmt}, fmt, {"mode": "metadata_only"})
        try:
            for ws in wb.worksheets:
                raw_tokens = _xlsx_raw_cells(wb, ws, limit + 1)
                rows = ws.iter_rows(values_only=False)
                header = next(rows, ())
                names = [cell.value for cell in header]
                for ci, name in enumerate(names):
                    _fact(bag, f"sheet:{ws.title}:column:{ci+1}", "name", name, "structure", "observed", rel,
                          {"kind": "xlsx_cell", "sheet": ws.title, "cell": header[ci].coordinate}, name, {"mode": "full"})
                    _fact(bag, f"sheet:{ws.title}:column:{ci+1}", "position", ci + 1, "syntax", "observed", rel,
                          {"kind": "xlsx_cell", "sheet": ws.title, "cell": header[ci].coordinate}, ci + 1, {"mode": "full"})
                for ri, cells in enumerate(rows, 2):
                    if ri > limit + 1: break
                    for ci, cell in enumerate(cells):
                        if cell.value is not None:
                            decoded = _json_value(cell.value)
                            raw_cell = {**raw_tokens.get(cell.coordinate, {}), "decoded_value": decoded,
                                        "number_format": cell.number_format, "data_type": cell.data_type}
                            _fact(bag, f"sheet:{ws.title}:column:{ci+1}", "observed_decoded_value", decoded, "syntax", "observed", rel,
                                  {"kind": "xlsx_cell", "sheet": ws.title, "sheet_member": ws._worksheet_path,
                                   "cell": cell.coordinate, "number_format": cell.number_format,
                                   "data_type": cell.data_type}, raw_cell,
                                  {"mode": "sample", "sample_limit": limit})
        finally:
            wb.close()
        return fmt, issues
    if fmt == "arff":
        if path.stat().st_size > 8 * 1024 * 1024:
            return fmt, ["unsupported: ARFF exceeds 8 MiB bounded parser limit"]
        try:
            from scipy.io import arff
            records, metadata = arff.loadarff(path)
        except ImportError:
            return fmt, ["dependency unavailable: scipy is required for arff"]
        except Exception as exc:
            raise ValueError(f"malformed arff: {exc}") from exc
        _fact(bag, "dataset", "serialization_format", fmt, "encoding", "observed", rel,
              {"kind": "validated_serialization", "format": fmt}, fmt, {"mode": "full"})
        raw_data_rows = []
        in_data = False
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            for line_number, raw_line in enumerate(stream, 1):
                if not in_data:
                    if raw_line.strip().lower().startswith("@data"):
                        in_data = True
                    continue
                if raw_line.strip() and not raw_line.lstrip().startswith("%"):
                    raw_data_rows.append((line_number, raw_line.rstrip("\r\n")))
                    if len(raw_data_rows) >= limit: break
        for ci, name in enumerate(metadata.names()):
            declaration = metadata[name]
            _fact(bag, f"attribute:{name}", "position", ci + 1, "syntax", "declared", rel,
                  {"kind": "arff_attribute", "name": name, "ordinal": ci + 1}, ci + 1, {"mode": "full"})
            # ARFF attributes and nominal domain members are actual declarations.
            _fact(bag, f"attribute:{name}", "attribute_type", str(declaration[0]), "structure", "declared", rel,
                  {"kind": "arff_attribute", "name": name}, str(declaration[0]), {"mode": "full"})
            if isinstance(declaration[1], (list, tuple)):
                domain = _json_value(declaration[1])
                _fact(bag, f"attribute:{name}", "declared_enum", domain, "value", "declared", rel,
                      {"kind": "arff_attribute", "name": name, "domain": True}, domain, {"mode": "full"})
        for ri, record in enumerate(records[:limit], 1):
            raw_line_number, raw_line = raw_data_rows[ri - 1] if ri <= len(raw_data_rows) else (None, None)
            for name in metadata.names():
                value = record[name]
                value = _json_value(value)
                _fact(bag, f"attribute:{name}", "observed_decoded_value", value, "syntax", "observed", rel,
                      {"kind": "arff_record", "row": ri, "line": raw_line_number, "attribute": name},
                      {"decoded_value": value, "raw_row_lexeme": raw_line},
                      {"mode": "sample", "sample_limit": limit})
        return fmt, issues
    if fmt == "zarr":
        members = bag["sources"][0].get("members", [])
        markers = [m for m in members if Path(m["path"]).name in {".zgroup", ".zarray", ".zattrs", ".zmetadata"}]
        found_v2 = False
        for member in markers:
            member_path = path / member["path"]
            if member["bytes"] > 8 * 1024 * 1024:
                raise ValueError(f"Zarr metadata member exceeds 8 MiB limit: {member['path']}")
            try:
                doc = json.loads(member_path.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
                raise ValueError(f"malformed Zarr metadata {member['path']}: {exc}") from exc
            if member_path.name in {".zgroup", ".zarray"}:
                if not isinstance(doc, dict) or doc.get("zarr_format") != 2:
                    raise ValueError(f"invalid Zarr-v2 metadata document: {member['path']}")
                found_v2 = True
            rel_member = member["path"]
            subject = "/" + str(Path(rel_member).parent).replace("\\", "/").replace(".", "").strip("/")
            subject = subject.rstrip("/") or "/"
            if member_path.name == ".zarray":
                props = {"shape": ("structure", "shape"), "chunks": ("structure", "chunks"),
                         "dtype": ("encoding", "dtype"), "compressor": ("encoding", "compressor"),
                         "filters": ("encoding", "filters"), "order": ("syntax", "order"),
                         "dimension_separator": ("syntax", "dimension_separator")}
                for key, (category, predicate) in props.items():
                    if key in doc:
                        _fact(bag, subject, predicate, doc[key], category, "declared", rel,
                              {"kind": "zarr_v2_metadata", "member_path": rel_member, "json_pointer": "/" + key},
                              doc[key], {"mode": "metadata_only", "chunks_read": False})
            elif member_path.name == ".zgroup":
                _fact(bag, subject, "group", True, "structure", "declared", rel,
                      {"kind": "zarr_v2_metadata", "member_path": rel_member, "json_pointer": "/zarr_format"},
                      doc.get("zarr_format"), {"mode": "metadata_only", "chunks_read": False})
            elif member_path.name == ".zattrs":
                if not isinstance(doc, dict): raise ValueError(f"invalid Zarr attributes document: {rel_member}")
                for key, value in doc.items():
                    category = "value" if str(key).lower() in {"unit", "units", "flag_values", "flag_meanings", "standard_name", "codes", "enum"} else "structure"
                    _fact(bag, subject, "attribute", {"name": key, "value": value}, category, "declared", rel,
                          {"kind": "zarr_v2_metadata", "member_path": rel_member, "json_pointer": "/" + str(key)},
                          value, {"mode": "metadata_only", "chunks_read": False})
            else:
                # Consolidated metadata contains a mapping of store paths to
                # the same raw metadata documents; retain it as explicit evidence.
                if not isinstance(doc, dict) or doc.get("zarr_consolidated_format") != 1:
                    raise ValueError(f"invalid consolidated Zarr metadata: {rel_member}")
                metadata = doc.get("metadata")
                if not isinstance(metadata, dict): raise ValueError(f"missing consolidated metadata mapping: {rel_member}")
                if any(isinstance(v, dict) and v.get("zarr_format") == 2 for v in metadata.values()): found_v2 = True
                _fact(bag, "/", "consolidated_metadata_entries", len(metadata), "structure", "declared", rel,
                      {"kind": "zarr_v2_metadata", "member_path": rel_member, "json_pointer": "/metadata"},
                      len(metadata), {"mode": "metadata_only", "chunks_read": False})
        if not found_v2:
            raise ValueError("Zarr store metadata does not identify Zarr format version 2")
        _fact(bag, "dataset", "serialization_format", "zarr-v2", "encoding", "observed", rel,
              {"kind": "validated_serialization", "format": "zarr-v2"}, "zarr-v2",
              {"mode": "metadata_only", "chunks_read": False})
        return fmt, issues
    if fmt in {"hdf5", "netcdf", "parquet"}:
        # The registry's deterministic readers supply physical metadata. Each
        # claim is attached to the source location named by the format reader.
        try:
            from ..extractors.base import ExtractionRequest
            from ..extractors.registry import extract_path
            outcome = extract_path(ExtractionRequest(str(path), format_hint=fmt, sample_limit=limit))
        except Exception as exc:
            return fmt, [f"metadata parser unavailable: {exc}"]
        if outcome.schema is None:
            return fmt, [issue.message for issue in outcome.issues] or ["metadata parser produced no schema"]
        coverage = {"mode": "metadata_only", "format_reader_status": outcome.status}
        _fact(bag, "dataset", "serialization_format", fmt, "encoding", "observed", rel,
              {"kind": "validated_serialization", "format": fmt}, fmt, coverage)
        metadata_by_path = {}
        if fmt == "netcdf":
            metadata_by_path = {item.get("field_path"): item.get("attributes", {})
                                for item in outcome.schema.metadata.get("cf_analysis", {}).get("variables", [])}
        elif fmt == "parquet":
            metadata_by_path = {item.get("field_path"): item.get("metadata", {})
                                for item in outcome.schema.metadata.get("parquet_analysis", {}).get("arrow_fields", [])}
        elif fmt == "hdf5":
            metadata_by_path = {item.get("path"): item.get("attributes", {}) for item in outcome.schema.groups}
        try:
            if fmt == "hdf5":
                import h5py
                with h5py.File(path, "r") as handle:
                    def read_attrs(name, node):
                        metadata_by_path["/" + name] = {str(k): _json_value(v) for k, v in node.attrs.items()}
                    handle.visititems(read_attrs)
            elif fmt == "netcdf":
                from scipy.io import netcdf_file
                with netcdf_file(path, "r", mmap=False) as handle:
                    for name, var in handle.variables.items():
                        metadata_by_path["/" + name] = {str(k): _json_value(v) for k, v in getattr(var, "_attributes", {}).items()}
        except ImportError:
            pass
        for i, field in enumerate(outcome.schema.fields):
            loc = field.field_path or field.field_name
            field_coverage = {**coverage, "source_locator": loc}
            _fact(bag, str(loc), "physical_type", field.physical_type, "structure", "declared", rel,
                  {"kind": "format_field_metadata", "format": fmt, "path": loc, "metadata_index": i, "property": "physical_type"}, field.physical_type, field_coverage)
            if field.shape is not None:
                _fact(bag, str(loc), "shape", field.shape, "structure", "declared", rel,
                      {"kind": "format_field_metadata", "format": fmt, "path": loc, "property": "shape"}, field.shape, coverage)
            attrs = metadata_by_path.get(loc, {}) or {}
            unit_value = attrs.get("units") or attrs.get("unit")
            if unit_value is not None:
                attr_key = "units" if fmt in {"netcdf", "parquet"} else ("units" if "units" in attrs else "unit")
                _fact(bag, str(loc), "declared_unit", unit_value, "value", "declared", rel,
                      {"kind": "format_attribute", "format": fmt, "path": loc, "attribute": attr_key}, unit_value, field_coverage)
            flag_values = attrs.get("flag_values")
            flag_meanings = attrs.get("flag_meanings")
            if flag_values is not None and flag_meanings is not None:
                meanings = str(flag_meanings).split()
                codes = list(flag_values) if isinstance(flag_values, (list, tuple)) else [flag_values]
                if len(codes) == len(meanings):
                    mapping = list(zip(codes, meanings))
                    _fact(bag, str(loc), "declared_code_meanings", mapping, "value", "declared", rel,
                          {"kind": "format_attribute", "format": fmt, "path": loc, "attributes": ["flag_values", "flag_meanings"]},
                          {"flag_values": flag_values, "flag_meanings": flag_meanings}, field_coverage)
        return fmt, issues
    return fmt, [f"format {fmt} is not currently supported by the adapter"]


def _parse_sidecars(sidecars: list[tuple[Path, str]], bag: dict) -> list[str]:
    issues = []
    for sidecar, rel in sidecars:
        if sidecar.suffix.lower() != ".json":
            issues.append(f"unsupported_sidecar_dialect: {rel} (expected .json JSON Schema)")
            continue
        try:
            schema = json.loads(sidecar.read_text(encoding="utf-8-sig"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError(f"malformed JSON Schema sidecar {rel}: {exc}") from exc
        props = schema.get("properties") if isinstance(schema, dict) else None
        if not isinstance(props, dict) or not ("$schema" in schema or schema.get("type") == "object"):
            issues.append(f"unsupported_sidecar_dialect: {rel}")
            continue
        required = schema.get("required", [])
        if not isinstance(required, list):
            raise ValueError(f"invalid JSON Schema required declaration in {rel}")
        required = set(required)
        for property_index, (name, declaration) in enumerate(props.items(), 1):
            if not isinstance(declaration, dict):
                continue
            subject = f"{name}"
            pointer = "/properties/" + str(name).replace("~", "~0").replace("/", "~1")
            coverage = {"mode": "full", "source_role": "sidecar", "bound_to_dataset_id": bag["dataset_id"]}
            common = {"kind": "json_schema_declaration", "json_pointer": pointer,
                      "sidecar_path": rel, "bound_dataset_id": bag["dataset_id"]}
            _fact(bag, subject, "property_lexical_position", property_index, "syntax", "declared", rel,
                  {**common, "json_pointer": "/properties"}, name, coverage)
            if "type" in declaration:
                _fact(bag, subject, "declared_type", declaration["type"], "structure", "declared", rel,
                      {**common, "json_pointer": pointer + "/type"}, declaration["type"], coverage)
            _fact(bag, subject, "required", name in required, "structure", "declared", rel,
                  {**common, "json_pointer": "/required"}, schema.get("required", []), coverage)
            for key, category, predicate in (("enum", "value", "declared_enum"),
                                             ("pattern", "syntax", "pattern"),
                                             ("format", "syntax", "declared_format"),
                                             ("minimum", "structure", "minimum"),
                                             ("maximum", "structure", "maximum")):
                if key in declaration:
                    _fact(bag, subject, predicate, declaration[key], category, "declared", rel,
                          {**common, "json_pointer": pointer + "/" + key}, declaration[key], coverage)
    return issues


def parse_dataset(path, *, root=None, format_hint=None, sample_limit=200, taxonomy=None, sidecars=())->dict:
    source = Path(path).resolve()
    base = Path(root).resolve() if root is not None else source.parent
    tax = _taxonomy_data(taxonomy)
    rel = source.name
    fmt = (format_hint or source.suffix.lower().lstrip("."))
    contained = False
    try:
        normalized_limit = int(sample_limit)
    except (TypeError, ValueError):
        normalized_limit = 0
    limit_error = "sample_limit must be a positive integer" if normalized_limit < 1 else None
    try:
        if limit_error: raise ValueError(limit_error)
        rel = _relative(source, base)
        contained = True
        ext = source.suffix.lower()
        fmt = (format_hint or ({".csv":"csv", ".tsv":"tsv", ".json":"json", ".jsonl":"jsonl", ".ndjson":"jsonl", ".xml":"xml", ".xsd":"xsd", ".arff":"arff", ".xlsx":"xlsx", ".h5":"hdf5", ".hdf5":"hdf5", ".nc":"netcdf", ".parquet":"parquet", ".zarr":"zarr"}.get(ext, "unknown"))).lower()
        if source.is_dir() and fmt == "zarr":
            members, raw = _zarr_inventory(source)
            raw_size = sum(m["bytes"] for m in members)
        elif source.is_file():
            source_hash, raw_size = _hash_file(source)
            raw = bytes.fromhex(source_hash)
            members = None
        else:
            raise ValueError("dataset path must be a regular file or Zarr-v2 directory store")
        content_hash = _digest(raw) if source.is_dir() else source_hash
        dataset_source = {"path": rel, "sha256": content_hash, "bytes": raw_size, "role": "dataset",
                          "hash_mode": "canonical_json_metadata_inventory" if source.is_dir() else "file_bytes"}
        if members is not None: dataset_source["members"] = members
        sources = [dataset_source]
        sidecar_paths=[]
        sidecar_items=[]
        for item in sidecars:
            sp = Path(item).resolve()
            sr = _relative(sp, base)
            if not sp.is_file(): raise ValueError(f"sidecar is not a file: {sp}")
            side_hash, side_size = _hash_file(sp)
            if side_size > 8 * 1024 * 1024:
                raise ValueError(f"unsupported: JSON Schema sidecar exceeds 8 MiB limit: {sr}")
            sources.append({"path": sr, "sha256": side_hash, "bytes": side_size, "role": "sidecar", "hash_mode": "file_bytes"})
            sidecar_paths.append(sr)
            sidecar_items.append((sp, sr))
        bundle = {"schema_version": SCHEMA_VERSION,
                  "dataset_id": "ds_" + _digest((rel + "\0" + content_hash).encode())[:24],
                  "taxonomy_sha256": _digest(_canonical(tax)), "sources": sources,
                  "parser": {"id": "four_category_evidence_adapter", "version": "1.0.0", "format": fmt,
                             "sample_limit": normalized_limit, "sidecar_paths": sidecar_paths,
                             "dependency_versions": _dependency_versions(fmt),
                             "limits": {"csv_sniff_bytes": 8192, "json_document_bytes": 8 * 1024 * 1024,
                                        "jsonl_line_bytes": 1024 * 1024, "arff_bytes": 8 * 1024 * 1024,
                                        "sidecar_bytes": 8 * 1024 * 1024, "zarr_metadata_member_bytes": 8 * 1024 * 1024,
                                        "xml_depth": 1024, "json_depth": 32, "json_nodes": 10000,
                                        "chunks_read": False if fmt == "zarr" else None}},
                  "facts": [], "evidence": [], "status": "pass", "issues": []}
        parsed_fmt, issues = _parse(source, rel, fmt, normalized_limit, bundle)
        issues.extend(_parse_sidecars(sidecar_items, bundle))
        if issues:
            bundle["issues"].extend(issues)
            bundle["status"] = "unsupported" if not bundle["facts"] else "partial"
        bundle["parser"]["format"] = parsed_fmt
    except Exception as exc:
        # Stable, replayable failure object still binds bytes if possible.
        members = None
        if not contained:
            digest = None
            raw_size = 0
        elif source.is_file():
            digest, raw_size = _hash_file(source)
            members = None
        elif contained and source.is_dir() and fmt == "zarr":
            try:
                members, raw_error = _zarr_inventory(source)
                digest = _digest(raw_error)
                raw_size = sum(m["bytes"] for m in members)
            except Exception:
                digest = None
                raw_size = 0
        else:
            digest = None
            raw_size = 0
        failure_sources = ([{"path": rel, "sha256": digest, "bytes": raw_size, "role": "dataset",
                             "hash_mode": "canonical_json_metadata_inventory" if members is not None else "file_bytes",
                             **({"members": members} if members is not None else {})}] if digest else [])
        if contained:
            for item in sidecars:
                sp = Path(item).resolve()
                try:
                    sr = _relative(sp, base)
                    if sp.is_file():
                        side_hash, side_size = _hash_file(sp)
                        failure_sources.append({"path": sr, "sha256": side_hash, "bytes": side_size, "role": "sidecar", "hash_mode": "file_bytes"})
                except ValueError:
                    continue
        bundle = {"schema_version": SCHEMA_VERSION,
                  "dataset_id": "ds_" + _digest((rel + "\0" + (digest or "unavailable")).encode())[:24],
                  "taxonomy_sha256": _digest(_canonical(tax)),
                  "sources": failure_sources,
                  "parser": {"id": "four_category_evidence_adapter", "version": "1.0.0", "format": fmt, "sample_limit": normalized_limit,
                             "sidecar_paths": [s["path"] for s in failure_sources if s["role"] == "sidecar"],
                             "dependency_versions": _dependency_versions(fmt),
                             "limits": {"csv_sniff_bytes": 8192, "json_document_bytes": 8 * 1024 * 1024,
                                        "jsonl_line_bytes": 1024 * 1024, "arff_bytes": 8 * 1024 * 1024,
                                        "sidecar_bytes": 8 * 1024 * 1024, "zarr_metadata_member_bytes": 8 * 1024 * 1024,
                                        "xml_depth": 1024, "json_depth": 32, "json_nodes": 10000,
                                        "chunks_read": False if fmt == "zarr" else None}},
                  "facts": [], "evidence": [], "status": "failed", "issues": [str(exc)]}
    bundle["bundle_sha256"] = _digest(_canonical(bundle))
    return bundle


def verify_dataset(bundle, source_root, *, taxonomy=None)->list[str]:
    errors=[]
    if not isinstance(bundle, dict): return ["bundle must be an object"]
    claimed = bundle.get("bundle_sha256")
    body = {k:v for k,v in bundle.items() if k != "bundle_sha256"}
    if claimed != _digest(_canonical(body)): errors.append("bundle_sha256 mismatch")
    source_root = Path(source_root).resolve()
    srcs = bundle.get("sources", [])
    if not srcs: errors.append("bundle has no sources"); return errors
    dataset = next((s for s in srcs if s.get("role") == "dataset"), None)
    if dataset is None: errors.append("bundle has no dataset source"); return errors
    for source in srcs:
        rel = source.get("path")
        path_obj = Path(str(rel))
        if not isinstance(rel, str) or path_obj.is_absolute() or PureWindowsPath(str(rel)).is_absolute() or ".." in path_obj.parts or "\\" in str(rel):
            errors.append(f"source path escapes source_root: {rel}")
            return errors
        for member in source.get("members", []):
            member_path = Path(str(member.get("path", "")))
            if member_path.is_absolute() or PureWindowsPath(str(member.get("path", ""))).is_absolute() or ".." in member_path.parts or "\\" in str(member.get("path", "")):
                errors.append(f"source member path escapes source_root: {member.get('path')}")
                return errors
    root_tax = _taxonomy_data(taxonomy)
    if bundle.get("taxonomy_sha256") != _digest(_canonical(root_tax)): errors.append("taxonomy hash mismatch")
    try:
        dataset_path = (source_root / dataset["path"]).resolve()
        opts = bundle.get("parser", {})
        sidecars = [source_root / s["path"] for s in srcs if s.get("role") == "sidecar"]
        replay = parse_dataset(dataset_path, root=source_root, format_hint=opts.get("format"),
                               sample_limit=opts.get("sample_limit", 200), taxonomy=root_tax, sidecars=sidecars)
        if _canonical({k:v for k,v in replay.items() if k != "bundle_sha256"}) != _canonical(body):
            errors.append("derived bundle differs from replayed source")
    except Exception as exc:
        errors.append(f"replay failed: {exc}")
    for fact in bundle.get("facts", []):
        for eid in fact.get("evidence_ids", []):
            evidence = next((e for e in bundle.get("evidence", []) if e.get("evidence_id") == eid), None)
            if evidence is None: errors.append(f"fact {fact.get('fact_id')} references missing evidence {eid}")
            elif evidence.get("source_path") not in {s.get("path") for s in srcs}:
                errors.append(f"evidence {eid} references unknown source")
    return errors
